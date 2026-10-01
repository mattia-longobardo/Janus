from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models import Device

COLUMNS = ["Name", "Status", "Group", "Access", "Static IP", "Current IP", "MAC", "Private MAC", "Vendor",
           "DHCP hostname", "DNS name", "First seen", "Last seen", "Last scan"]
ACCESS = {"authorized": "Full network", "lan_only": "LAN only", "pending": "Pending", "blocked": "Blocked"}
WIDTHS = [26, 10, 18, 14, 15, 15, 20, 12, 30, 22, 22, 18, 18, 18]


def _local(value: datetime | None, tz: ZoneInfo) -> datetime | None:
    return value.astimezone(tz).replace(tzinfo=None) if value is not None else None


def devices_workbook(devices: list[Device], tz: ZoneInfo) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Devices"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for d in devices:
        sheet.append([
            d.name,
            "Online" if d.online else "Offline",
            d.group.name if d.group else None,
            ACCESS.get(d.access.value, d.access.value),
            d.static_ip,
            d.last_ip,
            d.mac,
            "Yes" if d.private_mac else "No",
            d.vendor,
            d.dhcp_hostname,
            d.hostname,
            _local(d.first_seen, tz),
            _local(d.last_seen, tz),
            _local(d.last_scan_at, tz),
        ])
    for row in sheet.iter_rows(min_row=2, min_col=12, max_col=14):
        for cell in row:
            cell.number_format = "dd/mm/yyyy hh:mm"
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
