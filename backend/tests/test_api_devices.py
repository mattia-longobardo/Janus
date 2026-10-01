import uuid

import pytest
from sqlalchemy import select

from app.models import Access, Device, Event, Group


@pytest.fixture
def seeded(db):
    people = Group(name="People", color="#6FB7FF", icon="device", range_start="192.168.1.10",
                   range_end="192.168.1.19", default_access=Access.authorized)
    iot = Group(name="Power meters", color="#A6D86A", icon="device", range_start="192.168.1.120",
                range_end="192.168.1.129", default_access=Access.lan_only)
    db.add_all([people, iot])
    db.flush()
    laptop = Device(mac="00:00:5E:00:53:10", name="LAPTOP_A", hostname="laptop-a", group=people,
                    static_ip="192.168.1.10", access=Access.authorized)
    phone = Device(mac="02:00:5E:00:53:11", name="PHONE_A", hostname="phone-a", group=people,
                   static_ip="192.168.1.11", access=Access.authorized, private_mac=True)
    db.add_all([laptop, phone])
    db.flush()
    return {"people": people, "iot": iot, "laptop": laptop, "phone": phone}


def test_list_and_filter(client, seeded):
    all_devices = client.get("/api/devices").json()
    assert [d["name"] for d in all_devices] == ["LAPTOP_A", "PHONE_A"]
    assert all_devices[1]["private_mac"] is True
    assert client.get("/api/devices", params={"group_id": seeded["iot"].id}).json() == []
    assert client.get("/api/devices", params={"access": "authorized"}).status_code == 200


def test_get_unknown_device_is_404(client, seeded):
    response = client.get(f"/api/devices/{uuid.uuid4()}")
    assert response.status_code == 404
    assert response.json() == {"detail": "device not found"}


def test_rename_updates_hostname_and_logs(client, seeded, db):
    resp = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"name": "Laptop Work"})
    assert resp.status_code == 200
    assert resp.json()["hostname"] == "laptop-work"
    event = db.scalar(select(Event).where(Event.type == "device.updated"))
    assert event.mac == "00:00:5E:00:53:10"
    assert event.payload["changes"]["name"] == ["LAPTOP_A", "Laptop Work"]


def test_patch_device_rejects_ip_outside_group(client, seeded):
    resp = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"static_ip": "192.168.1.25"})
    assert resp.status_code == 422 and "outside the group range" in resp.json()["detail"]


def test_patch_rejects_taken_ip(client, seeded):
    taken = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"static_ip": "192.168.1.11"})
    assert taken.status_code == 422 and "already assigned" in taken.json()["detail"]


def test_moving_group_requires_ip_in_new_range(client, seeded):
    iot = seeded["iot"].id
    bad = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"group_id": iot})
    assert bad.status_code == 422 and "outside the group range" in bad.json()["detail"]
    ok = client.patch(f"/api/devices/{seeded['laptop'].id}",
                      json={"group_id": iot, "static_ip": "192.168.1.120", "access": "lan_only"})
    assert ok.status_code == 200
    assert (ok.json()["group_id"], ok.json()["static_ip"], ok.json()["access"]) == (iot, "192.168.1.120", "lan_only")


def test_keeping_own_ip_is_not_a_conflict(client, seeded):
    resp = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"static_ip": "192.168.1.10"})
    assert resp.status_code == 200


def test_unknown_group_is_422(client, seeded):
    resp = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"group_id": 9999})
    assert resp.status_code == 422


def test_patch_device_null_name_is_422(client, seeded):
    response = client.patch(f"/api/devices/{seeded['laptop'].id}", json={"name": None})
    assert response.status_code == 422


def _pending(db):
    device = Device(mac="00:00:5E:00:53:40", name="Unknown", hostname="unknown", group_id=None,
                    static_ip=None, access=Access.pending, last_ip="192.168.1.243")
    db.add(device)
    db.flush()
    return device


def test_pending_device_lists_without_group(client, db, seeded):
    _pending(db)
    body = client.get("/api/devices", params={"access": "pending"}).json()
    assert [(d["name"], d["group_id"], d["last_ip"]) for d in body] == [("Unknown", None, "192.168.1.243")]


def test_patch_ip_on_device_without_group_is_422(client, db, seeded):
    device = _pending(db)
    response = client.patch(f"/api/devices/{device.id}", json={"static_ip": "192.168.1.15"})
    assert response.status_code == 422 and "no group" in response.json()["detail"]


def test_patch_cannot_authorize_a_pending_device(client, db, seeded):
    device = _pending(db)
    response = client.patch(f"/api/devices/{device.id}", json={"access": "authorized"})
    assert response.status_code == 422 and "approve" in response.json()["detail"]


def test_rename_placeholder_without_mac_still_works(client, db, seeded):
    knob = Device(mac=None, name="KNOB", hostname="knob", group=seeded["people"], static_ip="192.168.1.14",
                  access=Access.authorized)
    db.add(knob)
    db.flush()
    assert client.patch(f"/api/devices/{knob.id}", json={"name": "KNOB 2"}).status_code == 200


def test_delete_device_forgets_it_and_logs(client, seeded, db):
    laptop = seeded["laptop"]
    response = client.delete(f"/api/devices/{laptop.id}")
    assert response.status_code == 204
    db.expire_all()
    assert db.get(Device, laptop.id) is None
    event = db.scalars(select(Event).where(Event.type == "device.deleted")).one()
    assert event.mac == "00:00:5E:00:53:10"
    assert event.payload == {"name": "LAPTOP_A", "ip": "192.168.1.10", "group": "People"}
    assert client.delete(f"/api/devices/{laptop.id}").status_code == 404


def test_export_devices_as_excel(client, seeded):
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get("/api/devices/export.xlsx", params={"group_id": seeded["people"].id})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="janus-devices-' in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:5] == ("Name", "Status", "Group", "Access", "Static IP")
    assert [r[0] for r in rows[1:]] == ["LAPTOP_A", "PHONE_A"]
    assert rows[2][6:8] == ("02:00:5E:00:53:11", "Yes")
    assert sheet.freeze_panes == "A2"
